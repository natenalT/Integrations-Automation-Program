import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

# =====================
# 📁 Set up file paths
# =====================
# These are the locations of your input CSV file, your crosswalk Excel file (used for mapping),
# and the final Excel file that will be created.
input_csv_path = r'C:\Users\joc4126\Desktop\CSV_File\Unfiltered\COCUH_HCA_Florida_Woodmont_Hospital.csv'
crosswalk_path = r'C:\Users\joc4126\Desktop\CSV_File\Data_Type_updated.xlsx'
output_excel_path = r'C:\Users\joc4126\Desktop\CSV_File\Filtered\HCA_Florida_Woodmont_Hospital.xlsx'

# Make sure the folder where we want to save the final Excel file exists.
# If it doesn't, Python will create it for us.
os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)

# =====================
# 📖 Load crosswalk file
# =====================
# This Excel file contains a lookup table that helps us translate codes into readable names.
# We only need the first 3 columns: one for the key, and two for the values we want to map.
crosswalk_df = pd.read_excel(crosswalk_path, usecols="A:C")

# =====================
# 🧹 Read CSV and remove duplicates
# =====================
# Load the raw CSV file into a DataFrame (like a spreadsheet in Python).
df = pd.read_csv(input_csv_path)

# Choose a few columns that help us identify duplicate rows.
# If two rows have the same values in these columns, we’ll keep only one.
cols_for_dedup = [df.columns[2], df.columns[3], df.columns[4],
                  df.columns[6], df.columns[9], df.columns[19]]

# Remove duplicate rows based on those columns.
df = df.drop_duplicates(subset=cols_for_dedup)

# =====================
# 🔄 Create mapping dictionaries
# =====================
# These dictionaries will help us convert codes (like "ADT_A01") into readable names.
message_type_map = dict(zip(crosswalk_df.iloc[:, 0], crosswalk_df.iloc[:, 1]))
data_type_map = dict(zip(crosswalk_df.iloc[:, 0], crosswalk_df.iloc[:, 2]))

# =====================
# 🧠 Define helper functions
# =====================
# These functions help us calculate new values based on each row of data.

# Figure out what kind of connection is used based on product names.
def get_connection_type(row):
    combined = f"{row['ob_product_name']} {row['ib_product_name']}".lower()
    if "kafka" in combined:
        return "Kafka"
    elif "waterpark" in combined:
        return "Waterpark"
    else:
        return "Cloverleaf"

# Decide if the data is going into HCA (Inbound) or out of HCA (Outbound).
def get_data_flow(row):
    if row['ib_product_name'] == 'iPeople' and row['ob_product_name'] == 'HCA 5.6':
        return 'Inbound'
    elif row['ib_product_name'] == 'HCA 5.6':
        return 'Outbound'
    elif row['ib_product_name'] == 'iPeople':
        return 'Outbound'
    else:
        return 'Inbound'

# Build a readable name for the interface, depending on direction and vendor.
def get_interface(row):
    data_type = data_type_map.get(row['ib_datatype'], '')
    direction = get_data_flow(row)
    if direction == 'Outbound':
        return f"{row['ob_vendor_name']} ({row['Vendor']}) - {data_type}"
    else:
        return f"{row['ib_vendor_name']} ({row['Vendor']}) - {data_type}"

# Choose which vendor name to show based on direction.
def get_product(row):
    direction = get_data_flow(row)
    return row['ob_vendor_name'] if direction == 'Outbound' else row['ib_vendor_name']

# =====================
# 🧱 Build a new DataFrame (like a table)
# =====================
# Create a new table with the columns we want to keep and display.
filtered_df = pd.DataFrame({
    'Product': df.apply(get_product, axis=1),
    'Interface': df.apply(get_interface, axis=1),
    'Facility': 'HCA Florida Woodmont Hospital',  # Hardcoded name for this file
    'Connection Type': df.apply(get_connection_type, axis=1),
    'Data Flow Direction': df.apply(get_data_flow, axis=1),
    'Message Type': df['ib_datatype'].map(message_type_map),
    'Data Type': df['ib_datatype'].map(data_type_map)
}).dropna(subset=['Product', 'Interface'])  # Remove rows missing key info

# =====================
# 🔠 Sort by Product name (A-Z)
# =====================
# This makes the final Excel file easier to read.
filtered_df = filtered_df.sort_values(by=['Product', 'Interface'], ascending=[True, True])

# =====================
# 🧩 Add vendor header rows
# =====================
# We want to group rows by vendor, and add a bold header row before each group.
output_rows = []
for product, group in filtered_df.groupby('Product', sort=False):
    # Create a bold header row with the vendor name
    vendor_row = {
        'Product': f"{product}",
        'Interface': '',
        'Facility': group['Facility'].iloc[0],
        'Connection Type': group['Connection Type'].iloc[0],
        'Data Flow Direction': '',
        'Message Type': '',
        'Data Type': ''
    }
    output_rows.append(vendor_row)
    
    # Add the actual data rows, but leave the Product column blank so it looks grouped
    for _, row in group.iterrows():
        row_copy = row.copy()
        row_copy['Product'] = ''
        output_rows.append(row_copy)

# Final table ready to export
final_df = pd.DataFrame(output_rows)

# =====================
# 💾 Save to Excel
# =====================
# Write the final table to an Excel file.
final_df.to_excel(output_excel_path, index=False)

# =====================
# 🎨 Format Excel file (make headers bold & centered)
# =====================
# Open the Excel file and make the vendor header rows bold and centered.
wb = load_workbook(output_excel_path)
ws = wb.active

bold_font = Font(bold=True)
center_align = Alignment(horizontal='center')

# Loop through each row and check if it's a header (Interface column is blank).
# If it is, make the whole row bold and centered.
for row in range(2, ws.max_row + 1):
    if ws.cell(row, 2).value == '':
        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).font = bold_font
            ws.cell(row, col).alignment = center_align

# Save the formatted Excel file.
wb.save(output_excel_path)
print("✅ Excel created, sorted by Product alphabetically, with vendor header rows.")
